import io
import datetime
from typing import List, Dict, Any, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


def format_seconds_to_time(seconds: float) -> str:
    """Format seconds into MM:SS or HH:MM:SS."""
    if seconds is None or seconds < 0:
        return "00:00"
    s = int(round(seconds))
    hrs = s // 3600
    mins = (s % 3600) // 60
    secs = s % 60
    if hrs > 0:
        return f"{hrs:02d}:{mins:02d}:{secs:02d}"
    return f"{mins:02d}:{secs:02d}"


def set_cell_background(cell, fill_hex: str):
    """Sets background color of a docx table cell."""
    tc_pr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding (in twips) for a cell."""
    tc_pr = cell._element.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def build_transcript_docx(
    title: str,
    duration: float,
    segments: List[Dict[str, Any]],
    script_text: Optional[str] = None,
    original_filename: Optional[str] = None,
    spaced_duration: Optional[float] = None
) -> io.BytesIO:
    """
    Generates a professionally styled Microsoft Word (.docx) document
    containing full transcript text and timestamped spoken phrases.
    """
    doc = Document()

    # Configure standard page margins (0.75 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run(title or "Meditation Audio Transcript")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)  # Deep navy

    # Subtitle / Category
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Audio Transcript & Spoken Cue Breakdown")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    # Metadata Summary Box (Table)
    meta_table = doc.add_table(rows=2, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    col_widths = [Inches(2.3), Inches(2.3), Inches(2.3)]
    for row in meta_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    dur_str = format_seconds_to_time(duration)
    if spaced_duration and spaced_duration > 0:
        dur_str += f" (Paced: {format_seconds_to_time(spaced_duration)})"

    num_phrases = len([s for s in segments if s.get("text")])
    date_str = datetime.datetime.utcnow().strftime("%B %d, %Y")

    meta_items = [
        ("Original File", original_filename or title or "Audio Track"),
        ("Audio Duration", dur_str),
        ("Total Phrases", f"{num_phrases} cues"),
        ("Export Date", date_str),
        ("Status", "AI Transcribed & Synchronized"),
        ("Pacing Mode", "Spaced Master" if spaced_duration else "Natural / Raw")
    ]

    for idx, (label, val) in enumerate(meta_items):
        r = idx // 3
        c = idx % 3
        cell = meta_table.cell(r, c)
        set_cell_background(cell, "F7FAFC")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

        lbl_run = p.add_run(f"{label}: ")
        lbl_run.font.name = "Calibri"
        lbl_run.font.size = Pt(9.5)
        lbl_run.font.bold = True
        lbl_run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

        val_run = p.add_run(str(val))
        val_run.font.name = "Calibri"
        val_run.font.size = Pt(9.5)
        val_run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 1: Full Continuous Transcript
    # -------------------------------------------------------------
    sec1_h = doc.add_paragraph()
    sec1_h.paragraph_format.space_before = Pt(10)
    sec1_h.paragraph_format.space_after = Pt(6)
    sec1_run = sec1_h.add_run("1. Full Continuous Transcript")
    sec1_run.font.name = "Calibri"
    sec1_run.font.size = Pt(14)
    sec1_run.font.bold = True
    sec1_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    # Determine full script text
    full_text = ""
    if script_text and script_text.strip():
        full_text = script_text.strip()
    elif segments:
        valid_texts = [s["text"].strip() for s in segments if s.get("text") and s["text"].strip()]
        full_text = "\n\n".join(valid_texts)

    if not full_text:
        full_text = "(No transcript text detected or recorded for this project.)"

    paragraphs = [p.strip() for p in full_text.split("\n\n") if p.strip()]
    for p_content in paragraphs:
        body_p = doc.add_paragraph()
        body_p.paragraph_format.space_before = Pt(0)
        body_p.paragraph_format.space_after = Pt(6)
        body_p.paragraph_format.line_spacing = 1.2
        r = body_p.add_run(p_content)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # -------------------------------------------------------------
    # SECTION 2: Timestamped Spoken Phrase Breakdown
    # -------------------------------------------------------------
    if segments and len(segments) > 0:
        sec2_h = doc.add_paragraph()
        sec2_h.paragraph_format.space_before = Pt(12)
        sec2_h.paragraph_format.space_after = Pt(6)
        sec2_run = sec2_h.add_run("2. Timestamped Phrase Breakdown & Cues")
        sec2_run.font.name = "Calibri"
        sec2_run.font.size = Pt(14)
        sec2_run.font.bold = True
        sec2_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        table_col_widths = [Inches(1.5), Inches(4.3), Inches(1.2)]
        for idx, width in enumerate(table_col_widths):
            table.rows[0].cells[idx].width = width

        # Header row
        hdr_cells = table.rows[0].cells
        hdr_titles = ["Time Interval", "Spoken Phrase / Meditation Cue", "Pause After"]
        for idx, h_text in enumerate(hdr_titles):
            cell = hdr_cells[idx]
            set_cell_background(cell, "1B365D")
            set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if idx == 0 or idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(h_text)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        for s_idx, seg in enumerate(segments):
            row_cells = table.add_row().cells
            for idx, width in enumerate(table_col_widths):
                row_cells[idx].width = width

            bg_hex = "F8F9FA" if s_idx % 2 == 1 else "FFFFFF"

            start_t = seg.get("start_time", 0.0)
            end_t = seg.get("end_time", 0.0)
            text = seg.get("text", "").strip() or "—"
            pause_dur = seg.get("pause_duration")
            if pause_dur is None:
                pause_dur = seg.get("natural_silence_dur", 0.0)

            time_str = f"[{format_seconds_to_time(start_t)} - {format_seconds_to_time(end_t)}]"
            pause_str = f"{pause_dur:.1f}s" if pause_dur else "—"

            for c_idx, cell in enumerate(row_cells):
                set_cell_background(cell, bg_hex)
                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)

                if c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(time_str)
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
                elif c_idx == 1:
                    run = p.add_run(text)
                    run.font.name = "Calibri"
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(pause_str)
                    run.font.name = "Calibri"
                    run.font.size = Pt(9.5)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0xD9, 0x77, 0x06)  # Warm amber

    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io
